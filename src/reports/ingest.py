"""CSV/XLSX ingestion with one fixed template: rows of `field,value[,description]`.

No column guessing: a mis-mapped column would silently produce a report with wrong figures.
"""
import csv
import io
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from pydantic import ValidationError

from .schemas import CashFlowMetrics

MAX_BYTES = 1_000_000
MAX_ROWS = 100
TEXT_FIELDS = ("company_name", "period", "language")
METRIC_FIELDS = tuple(CashFlowMetrics.model_fields)

TEMPLATE_ROWS = [
    ("field", "value", "description"),
    ("company_name", "", "Nombre de la empresa / Company name"),
    ("period", "Q2 2026", "Trimestre Q1-Q4 y año / Quarter and year"),
    ("language", "es", "es | en"),
    ("revenue", "", "Ingresos, millones USD / Revenue, USD millions"),
    ("ebitda", "", "EBITDA, millones USD"),
    ("operating_cash_flow", "", "Flujo de caja operativo / Operating cash flow"),
    ("capex", "", "CAPEX"),
    ("working_capital_change", "0", "Variación de capital de trabajo / Working capital change"),
    ("cash", "", "Caja / Cash"),
    ("net_debt", "", "Deuda neta / Net debt"),
]


class IngestError(ValueError):
    def __init__(self, errors):
        super().__init__("; ".join(errors))
        self.errors = errors


def template_csv():
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(TEMPLATE_ROWS)
    return out.getvalue()


def parse_number(text):
    """Parse '1,234.5', '1.234,5', '$ 842.3' or '-12'. The last separator is the decimal mark;
    a lone comma followed by exactly three digits ('1,234') is read as a thousands separator."""
    s = text.strip().replace("$", "").replace(" ", "").replace(" ", "")
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in s:
        whole, _, frac = s.rpartition(",")
        s = s.replace(",", "") if (s.count(",") > 1 or len(frac) == 3) else f"{whole}.{frac}"
    return Decimal(s)


def _rows_from_csv(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")  # Excel on Windows exports CSV as cp1252/latin-1
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") > first.count(",") else ","  # Spanish-locale Excel uses ';'
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))[:MAX_ROWS]


def _rows_from_xlsx(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.active.iter_rows(max_row=MAX_ROWS, max_col=2, values_only=True)]
    finally:
        wb.close()


def _to_decimal(value):
    if isinstance(value, bool):
        raise InvalidOperation
    if isinstance(value, (int, float)):
        # Excel stores doubles: 0.1 + 0.2 arrives as 0.30000000000000004, so round to cents.
        return Decimal(repr(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return parse_number(str(value))


def parse_upload(data):
    """Return {"fields": {company_name?, period?, language?}, "metrics": CashFlowMetrics}."""
    if len(data) > MAX_BYTES:
        raise IngestError([f"file larger than {MAX_BYTES // 1000} KB"])
    is_xlsx = data[:4] == b"PK\x03\x04"  # XLSX is a ZIP container
    try:
        rows = _rows_from_xlsx(data) if is_xlsx else _rows_from_csv(data)
    except Exception as e:  # corrupt or unsupported file
        raise IngestError([f"could not read {'XLSX' if is_xlsx else 'CSV'} file: {e}"]) from e

    errors, fields, metrics = [], {}, {}
    for n, row in enumerate(rows, start=1):
        key = str(row[0]).strip().lower() if row and row[0] is not None else ""
        value = row[1] if len(row) > 1 else None
        if not key or key == "field":
            continue
        if key in fields or key in metrics:
            errors.append(f"row {n}: duplicate field '{key}'")
        elif key in TEXT_FIELDS:
            if value is not None and str(value).strip():
                fields[key] = str(value).strip()
        elif key in METRIC_FIELDS:
            if value is None or not str(value).strip():
                continue  # blank: required fields are reported as missing by validation below
            try:
                metrics[key] = _to_decimal(value)
            except (InvalidOperation, ValueError):
                errors.append(f"row {n}: '{key}' is not a number: {value!r}")
        else:
            errors.append(f"row {n}: unknown field '{key}'")
    if errors:
        raise IngestError(errors)
    try:
        return {"fields": fields, "metrics": CashFlowMetrics(**metrics)}
    except ValidationError as e:
        raise IngestError([f"{'.'.join(map(str, err['loc'])) or 'metrics'}: {err['msg']}"
                           for err in e.errors()]) from e
